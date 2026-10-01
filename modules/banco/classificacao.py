"""Passo 5c: quem recebeu cada saída do extrato (e quem mandou cada entrada).

Origens, da mais forte para a mais fraca. Uma origem nunca é sobrescrita pelas de baixo:
1. `manual`    um humano decidiu. Nada automático mexe.
2. `sistema`   transferência entre contas próprias, provada pelo PAR: saída numa conta = entrada na
               outra, mesmo valor, até 2 dias depois, com o nome da própria empresa no texto da entrada.
               O texto sozinho não serve: "PIX RECEBIDO REM: RODIZA ALIMENTOS" também é venda repassada.
3. `planilha`  o que o usuário já classificou, linha a linha, na planilha 15_Bradesco_tratado.
4. `aprendida` regra de texto aprendida da planilha (LEFT 33 do histórico, a chave que o usuário já usa).

O texto do banco diz QUEM recebeu, nem sempre O QUE foi pago: o boleto "COBRANCA GIRAFFAS" é royalties,
fundo de marketing ou REFIN. Nesses casos a regra traz o favorecido e deixa a classificação vazia.
"""
import calendar
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from modules.banco.extratos import _data, _num
from modules.banco.models import GRUPOS_CATEGORIA, ContaBancaria, ExtratoCategoria, ExtratoLancamento, ExtratoRegra
from modules.pagamentos.models import Fornecedor
from modules.tenants.models import Tenant

TRANSFERENCIA = "TRANSFERENCIA ENTRE CONTAS"
JANELA_TRANSFERENCIA_DIAS = 2
# O Bradesco trocou o texto do Pix entre exportações; a chave usa sempre a forma nova.
_SINONIMOS_INICIO = (("TRANSFERENCIA PIX DES:", "PIX ENVIADO DES:"), ("TRANSFERENCIA PIX REM:", "PIX RECEBIDO REM:"))
_ENVIO = re.compile(r"PIX ENVIADO|TRANSFERENCIA|TED|DOC\b")
_DATA_NO_FIM = re.compile(r"\s+\d{2}/\d{2}(/\d{2,4})?$")


def normalizar(texto: str) -> str:
    t = " ".join(str(texto).upper().split())
    for velho, novo in _SINONIMOS_INICIO:
        if t.startswith(velho):
            t = novo + t[len(velho):]
    return t


def chave(historico: str) -> str:
    """LEFT 33 do histórico padronizado, sem a data que o banco põe no fim do Pix ('... 04/08')."""
    return _DATA_NO_FIM.sub("", normalizar(historico))[:33].rstrip()


def _celula(v) -> str:
    """Texto de célula da planilha: maiúsculo, sem pontuação solta ('TRANSF. PONTO' = 'TRANSF, PONTO')."""
    if v is None or str(v).strip().lower() in ("", "nan"):
        return ""
    return " ".join(re.sub(r"[.,]", " ", str(v).upper()).split())


@dataclass
class Dicionario:
    grafias: dict[str, str] = field(default_factory=dict)
    classificacoes: dict[str, str] = field(default_factory=dict)
    fornecedor_crm: dict[str, str] = field(default_factory=dict)
    sem_favorecido: set[str] = field(default_factory=set)
    regras_manuais: list[dict] = field(default_factory=list)

    @classmethod
    def ler(cls, caminho: Path | None) -> "Dicionario":
        if caminho is None:
            return cls()
        d = json.loads(caminho.read_text(encoding="utf-8"))
        return cls({_celula(k): _celula(v) for k, v in d.get("grafias", {}).items()},
                   {_celula(k): _celula(v) for k, v in d.get("classificacoes", {}).items()},
                   {_celula(k): v for k, v in d.get("fornecedor_crm", {}).items()},
                   {_celula(v) for v in d.get("sem_favorecido", [])}, d.get("regras_manuais", []))

    def favorecido(self, v) -> str | None:
        f = _celula(v)
        return None if not f or f in self.sem_favorecido else self.grafias.get(f, f)

    def classificacao(self, v) -> str | None:
        c = _celula(v)
        return self.classificacoes.get(c, c) or None


@dataclass
class LinhaPlanilha:
    data: date
    historico: str
    valor: Decimal
    favorecido: str | None
    classificacao: str | None


def ler_planilha(caminho: Path, dic: Dicionario) -> list[LinhaPlanilha]:
    """Aba com cabeçalho Data / Lançamento / Movimentação / Classificação / Fornecedor (a 'Sheet0' da tratada)."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            linhas = ws.iter_rows(values_only=True)
            for cab in linhas:
                nomes = [str(c).strip() if c is not None else "" for c in cab]
                if {"Data", "Lançamento", "Movimentação", "Classificação", "Fornecedor"} <= set(nomes):
                    break
            else:
                continue
            i = {n: nomes.index(n) for n in ("Data", "Lançamento", "Movimentação", "Classificação", "Fornecedor")}
            saida = []
            for r in linhas:
                d, v = _data(r[i["Data"]]), _num(r[i["Movimentação"]])
                if d is None or v is None or not r[i["Lançamento"]]:
                    continue
                saida.append(LinhaPlanilha(d, str(r[i["Lançamento"]]), v, dic.favorecido(r[i["Fornecedor"]]),
                                           dic.classificacao(r[i["Classificação"]])))
            return saida
    finally:
        wb.close()
    raise ValueError(f"{caminho.name}: nenhuma aba com Data/Lançamento/Movimentação/Classificação/Fornecedor")


def _fornecedores(db: Session, tenant_id: int, dic: Dicionario) -> tuple[dict[str, int], list[str]]:
    por_nome = {f.nome: f.id for f in db.scalars(select(Fornecedor).where(Fornecedor.tenant_id == tenant_id))}
    faltam = sorted(n for n in dic.fornecedor_crm.values() if n not in por_nome)
    return {fav: por_nome[n] for fav, n in dic.fornecedor_crm.items() if n in por_nome}, faltam


@dataclass
class ResultadoSemente:
    linhas_planilha: int = 0
    pareadas: int = 0
    sem_par: list[LinhaPlanilha] = field(default_factory=list)
    regras_novas: int = 0
    regras_atualizadas: int = 0
    regras_desligadas: int = 0
    regras_sem_classificacao: list[str] = field(default_factory=list)
    texto_diferente: list[tuple[LinhaPlanilha, str]] = field(default_factory=list)   # (linha, texto no banco)
    fornecedores_faltando: list[str] = field(default_factory=list)


def semear_da_planilha(db: Session, tenant_id: int, conta: ContaBancaria, linhas: list[LinhaPlanilha],
                       dic: Dicionario) -> ResultadoSemente:
    """Liga cada linha da planilha ao lançamento do extrato (data + valor; empate → mesma chave de texto),
    copia favorecido/classificação para ele e aprende uma regra por chave de texto."""
    res = ResultadoSemente(linhas_planilha=len(linhas))
    if not linhas:
        return res
    forn, res.fornecedores_faltando = _fornecedores(db, tenant_id, dic)
    ext = db.scalars(select(ExtratoLancamento).where(
        ExtratoLancamento.conta_id == conta.id, ExtratoLancamento.data >= min(l.data for l in linhas),
        ExtratoLancamento.data <= max(l.data for l in linhas))
        .order_by(ExtratoLancamento.data, ExtratoLancamento.ordem, ExtratoLancamento.id)).all()
    candidatos: dict[tuple, list[ExtratoLancamento]] = defaultdict(list)
    for l in ext:
        candidatos[(l.data, l.valor)].append(l)
    pares: list[tuple[LinhaPlanilha, ExtratoLancamento]] = []
    for p in linhas:
        livres = candidatos[(p.data, p.valor)]
        if not livres:
            res.sem_par.append(p)
            continue
        k = chave(p.historico)
        l = next((c for c in livres if chave(c.historico) == k), livres[0])
        livres.remove(l)
        pares.append((p, l))
    res.pareadas = len(pares)

    # Uma regra por (chave, natureza): favorecido mais frequente; classificação só se for sempre a mesma.
    # Só aprende com linha cujo texto na planilha é o do banco: exportações antigas cortavam o favorecido
    # ("PAGTO ELETRON COBRANCA") e a classificação dada às cegas não vale para o texto completo.
    exemplos: dict[tuple[str, str], list[LinhaPlanilha]] = defaultdict(list)
    for p, l in pares:
        if chave(p.historico) != chave(l.historico):
            res.texto_diferente.append((p, l.historico))
        elif p.favorecido:
            exemplos[(chave(l.historico), l.natureza)].append(p)
    existentes = {(r.padrao_texto, r.natureza): r for r in db.scalars(
        select(ExtratoRegra).where(ExtratoRegra.tenant_id == tenant_id, ExtratoRegra.tipo_match == "prefixo"))}
    regras: dict[tuple[str, str], ExtratoRegra] = {}
    for (k, nat), ps in sorted(exemplos.items()):
        fav = Counter(p.favorecido for p in ps).most_common(1)[0][0]
        classes = sorted({p.classificacao for p in ps if p.classificacao})
        cls = classes[0] if len(classes) == 1 and all(p.classificacao for p in ps) else None
        if cls is None:
            res.regras_sem_classificacao.append(f"{k} ({nat}): {fav}")
        r = existentes.get((k, nat))
        if r is None:
            r = ExtratoRegra(tenant_id=tenant_id, padrao_texto=k, tipo_match="prefixo", natureza=nat, origem="planilha")
            db.add(r)
            res.regras_novas += 1
        elif r.origem == "manual":          # humano manda: a planilha não mexe
            regras[(k, nat)] = r
            continue
        else:
            res.regras_atualizadas += 1
        r.favorecido, r.classificacao, r.fornecedor_id, r.ocorrencias, r.ativo = fav, cls, forn.get(fav), len(ps), True
        r.classificacoes = classes
        regras[(k, nat)] = r
    for k, r in existentes.items():          # regra da planilha que perdeu todo o respaldo: desliga
        if r.origem == "planilha" and k not in exemplos and r.ativo:
            r.ativo = False
            res.regras_desligadas += 1
    db.flush()

    for p, l in pares:
        if l.classificacao_origem in ("manual", "sistema") or not (p.favorecido or p.classificacao):
            continue
        r = regras.get((chave(l.historico), l.natureza))
        l.favorecido, l.classificacao, l.fornecedor_id = p.favorecido, p.classificacao, forn.get(p.favorecido)
        l.classificacao_origem, l.regra_id = "planilha", r.id if r else None
    db.flush()
    return res


def gravar_regras_manuais(db: Session, tenant_id: int, dic: Dicionario, usuario_id: int | None = None) -> int:
    """Regras decididas fora da planilha (ver o dicionário do cliente). Têm prioridade sobre as aprendidas."""
    forn, _ = _fornecedores(db, tenant_id, dic)
    n = 0
    for m in dic.regras_manuais:
        tipo = m.get("tipo", "prefixo")
        padrao = m["padrao"] if tipo == "regex" else normalizar(m["padrao"]) if tipo == "contem" else chave(m["padrao"])
        if tipo == "regex":
            re.compile(padrao)                       # regex inválida falha aqui, antes de gravar
        r = db.scalar(select(ExtratoRegra).where(ExtratoRegra.tenant_id == tenant_id, ExtratoRegra.padrao_texto == padrao,
                                                 ExtratoRegra.tipo_match == tipo, ExtratoRegra.natureza == m["natureza"]))
        if r is None:
            r = ExtratoRegra(tenant_id=tenant_id, padrao_texto=padrao, tipo_match=tipo, natureza=m["natureza"])
            db.add(r)
        fav = dic.favorecido(m["favorecido"])
        cls = dic.classificacao(m.get("classificacao"))
        r.favorecido, r.classificacao, r.classificacoes = fav, cls, [cls] if cls else []
        r.fornecedor_id, r.origem, r.ativo, r.atualizado_por = forn.get(fav), "manual", True, usuario_id
        r.parametros = m.get("parametros", {})
        n += 1
    db.flush()
    return n


def gravar_categorias(db: Session, tenant_id: int, caminho: Path, usuario_id: int | None = None) -> int:
    """Carrega/atualiza a tabela classe → categoria do DRE a partir do JSON do cliente (idempotente)."""
    itens = json.loads(caminho.read_text(encoding="utf-8"))["categorias"]
    atuais = {(c.natureza, c.classificacao, c.favorecido): c for c in db.scalars(
        select(ExtratoCategoria).where(ExtratoCategoria.tenant_id == tenant_id))}
    for natureza, classificacao, favorecido, categoria, grupo in itens:
        if natureza not in ("C", "D") or grupo not in GRUPOS_CATEGORIA:
            raise ValueError(f"categoria inválida: {natureza} {classificacao} {grupo}")
        k = (natureza, _celula(classificacao), _celula(favorecido) or None)
        c = atuais.get(k) or ExtratoCategoria(tenant_id=tenant_id, natureza=k[0], classificacao=k[1], favorecido=k[2])
        c.categoria, c.grupo, c.atualizado_por = categoria, grupo, usuario_id
        db.add(c)
    db.flush()
    return len(itens)


def _casa(r: ExtratoRegra, k: str, texto: str) -> bool:
    if r.tipo_match == "prefixo":
        return k.startswith(r.padrao_texto)
    if r.tipo_match == "contem":
        return normalizar(r.padrao_texto) in texto
    return re.search(r.padrao_texto, texto) is not None


def aplicar_regras(db: Session, tenant_id: int) -> int:
    """Classifica o que ainda não tem dono (ou foi aprendido antes) pela regra mais específica."""
    regras = sorted(db.scalars(select(ExtratoRegra).where(ExtratoRegra.tenant_id == tenant_id,
                                                          ExtratoRegra.ativo.is_(True))),
                    key=lambda r: (r.origem != "manual", -len(r.padrao_texto)))
    n = 0
    for l in db.scalars(select(ExtratoLancamento).where(
            ExtratoLancamento.tenant_id == tenant_id,
            or_(ExtratoLancamento.classificacao_origem.is_(None), ExtratoLancamento.classificacao_origem == "aprendida"))):
        texto, k = normalizar(l.historico), chave(l.historico)
        r = next((r for r in regras if r.natureza == l.natureza and _casa(r, k, texto)), None)
        if r is not None:
            l.favorecido, l.classificacao, l.fornecedor_id = r.favorecido, r.classificacao, r.fornecedor_id
            l.classificacao_origem, l.regra_id = "aprendida", r.id
            n += 1
        elif l.classificacao_origem == "aprendida":        # a regra foi desligada
            l.favorecido = l.classificacao = l.fornecedor_id = l.classificacao_origem = l.regra_id = None
    db.flush()
    return n


def _dia_util(d: date) -> date:
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def _vencimento(ano: int, mes: int, dia) -> date:
    ultimo = calendar.monthrange(ano, mes)[1]
    return date(ano, mes, ultimo if dia == "ultimo" else min(int(dia), ultimo))


def classe_pelo_calendario(data: date, valor: Decimal, itens: list[dict]) -> str | None:
    """Calendário de pagamento de um favorecido. Ex. (franqueadora): REFIN tem valores fixos; ROYALTIES vence
    dia 20 e FUNDO MKT no último dia do mês; em fim de semana/feriado paga no próximo dia útil. Valor fixo decide
    primeiro; senão, o vencimento (deste mês ou do anterior) mais perto da data do pagamento, dentro da janela
    [vencimento − antes, 1º dia útil a partir do vencimento + depois]. Empate ou nada na janela → None (humano)."""
    for i in itens:
        if abs(valor) in {Decimal(str(v)) for v in i.get("valores", [])}:
            return i["classificacao"]
    candidatos = []
    anterior = data.replace(day=1) - timedelta(days=1)
    for i in itens:
        if "dia" not in i:
            continue
        for ano, mes in ((data.year, data.month), (anterior.year, anterior.month)):
            venc = _vencimento(ano, mes, i["dia"])
            pago = _dia_util(venc)
            if venc - timedelta(days=i.get("antes", 3)) <= data <= pago + timedelta(days=i.get("depois", 4)):
                candidatos.append((abs((data - pago).days), i["classificacao"]))
    candidatos.sort()
    if not candidatos or (len(candidatos) > 1 and candidatos[0][0] == candidatos[1][0]
                          and candidatos[0][1] != candidatos[1][1]):
        return None
    return candidatos[0][1]


def aplicar_calendarios(db: Session, tenant_id: int) -> tuple[int, list[tuple[ExtratoLancamento, str]]]:
    """Regra com `parametros.calendario` decide a classe dos lançamentos do seu favorecido pela data/valor.
    Vale acima da planilha (é regra de negócio informada pelo usuário); `manual` continua intocável.
    Devolve (quantos decididos, [(lançamento, classe anterior)] onde a planilha/regra dizia outra coisa)."""
    n, trocas = 0, []
    for r in db.scalars(select(ExtratoRegra).where(ExtratoRegra.tenant_id == tenant_id, ExtratoRegra.ativo.is_(True))):
        itens = (r.parametros or {}).get("calendario")
        if not itens:
            continue
        for l in db.scalars(select(ExtratoLancamento).where(
                ExtratoLancamento.tenant_id == tenant_id, ExtratoLancamento.natureza == r.natureza,
                ExtratoLancamento.favorecido == r.favorecido,
                or_(ExtratoLancamento.classificacao_origem.is_(None), ExtratoLancamento.classificacao_origem != "manual"))):
            cls = classe_pelo_calendario(l.data, l.valor, itens)
            if cls is None:
                continue
            if l.classificacao and l.classificacao != cls:
                trocas.append((l, l.classificacao))
            l.classificacao, l.classificacao_origem, l.regra_id = cls, "sistema", r.id
            l.fornecedor_id = l.fornecedor_id or r.fornecedor_id
            n += 1
    db.flush()
    return n, sorted(trocas, key=lambda t: t[0].data)


def _marca_da_empresa(razao_social: str) -> str:
    """'Rodiza Alimentos Ltda' → 'RODIZA': a 1ª palavra com 4+ letras da razão social."""
    return next((p for p in normalizar(razao_social).split() if len(p) >= 4), normalizar(razao_social))


def marcar_transferencias(db: Session, tenant_id: int) -> int:
    """Pares saída ↔ entrada entre contas próprias. Só marca quando o par é único (senão fica para humano)."""
    contas = {c.id: c for c in db.scalars(select(ContaBancaria).where(ContaBancaria.tenant_id == tenant_id))}
    if len(contas) < 2:
        return 0
    marca = _marca_da_empresa(db.get(Tenant, tenant_id).razao_social)
    livres = [l for l in db.scalars(select(ExtratoLancamento).where(
        ExtratoLancamento.tenant_id == tenant_id,
        or_(ExtratoLancamento.classificacao_origem.is_(None), ExtratoLancamento.classificacao_origem != "manual"))
        .order_by(ExtratoLancamento.data, ExtratoLancamento.ordem, ExtratoLancamento.id))]
    entradas: dict[Decimal, list[ExtratoLancamento]] = defaultdict(list)
    for l in livres:
        if l.valor > 0 and marca in normalizar(l.historico):
            entradas[l.valor].append(l)
    usados, n = set(), 0
    for s in livres:
        if s.valor >= 0 or not _ENVIO.search(normalizar(s.historico)):
            continue
        cands = [e for e in entradas[-s.valor] if e.id not in usados and e.conta_id != s.conta_id
                 and 0 <= (e.data - s.data).days <= JANELA_TRANSFERENCIA_DIAS]
        if not cands:
            continue
        perto = min((e.data - s.data).days for e in cands)
        cands = [e for e in cands if (e.data - s.data).days == perto]
        if len(cands) != 1:
            continue
        e = cands[0]
        usados.add(e.id)
        for l, outra in ((s, e), (e, s)):
            l.favorecido, l.classificacao, l.fornecedor_id = contas[outra.conta_id].apelido, TRANSFERENCIA, None
            l.classificacao_origem, l.regra_id = "sistema", None
        n += 1
    db.flush()
    return n


@dataclass
class Resumo:
    regras: int
    calendario: int
    trocas_calendario: list[tuple[ExtratoLancamento, str]]
    transferencias: int

    def __str__(self) -> str:
        return (f"{self.regras} lançamento(s) classificados por regra · {self.calendario} pelo calendário de pagamento "
                f"({len(self.trocas_calendario)} diferente(s) do que estava) · "
                f"{self.transferencias} transferência(s) entre contas")


def classificar(db: Session, tenant_id: int) -> Resumo:
    """O que roda a cada importação: regras de texto → calendários de pagamento → transferências entre contas."""
    n_regras = aplicar_regras(db, tenant_id)
    n_cal, trocas = aplicar_calendarios(db, tenant_id)
    return Resumo(n_regras, n_cal, trocas, marcar_transferencias(db, tenant_id))


def cobertura(db: Session, tenant_id: int, de: date | None = None) -> list[dict]:
    """Por conta e natureza: quanto já tem favorecido e quanto falta (com as maiores chaves pendentes)."""
    q = select(ExtratoLancamento).where(ExtratoLancamento.tenant_id == tenant_id)
    if de:
        q = q.where(ExtratoLancamento.data >= de)
    contas = {c.id: c.apelido for c in db.scalars(select(ContaBancaria).where(ContaBancaria.tenant_id == tenant_id))}
    grupos: dict[tuple, dict] = {}
    for l in db.scalars(q):
        g = grupos.setdefault((contas[l.conta_id], l.natureza), {
            "conta": contas[l.conta_id], "natureza": l.natureza, "total": 0, "valor_total": Decimal(0),
            "por_origem": Counter(), "valor_pendente": Decimal(0), "pendentes": defaultdict(lambda: [0, Decimal(0)])})
        g["total"] += 1
        g["valor_total"] += l.valor
        g["por_origem"][l.classificacao_origem or "pendente"] += 1
        if l.favorecido is None:
            g["valor_pendente"] += l.valor
            p = g["pendentes"][chave(l.historico)]
            p[0] += 1
            p[1] += l.valor
    for g in grupos.values():
        g["pendentes"] = sorted(((k, q, v) for k, (q, v) in g["pendentes"].items()), key=lambda x: -abs(x[2]))
    return sorted(grupos.values(), key=lambda g: (g["conta"], g["natureza"]))
