"""Leitura dos extratos Bradesco e Santander: ENTRADAS e SAÍDAS, em ordem cronológica, com saldo.

Mesma lógica dos parsers da Gestão (Giraffas-Datalake/04_scripts/conciliacao_fontes/banco_*.py),
copiada e estendida SEM mudar as regras de leitura:
- Bradesco (.xlsx, .xls antigo ou .csv ';' latin1): cabeçalho Data / Lançamento / Dcto. / Crédito (R$) /
  Débito (R$) / Saldo (R$) algumas linhas abaixo do topo; a leitura PARA na 1ª linha "Total" (depois vêm
  "Últimos Lançamentos", repetidos em todo arquivo, e às vezes uma tabela de Aplicações).
- Santander (.xlsx): cabeçalho Data / Histórico / Documento / Valor (R$) / Saldo (R$); vem do mais NOVO
  para o mais antigo e cada exportação tem no máximo 440 linhas (corta os dias mais antigos).
- Santander internet banking ("EXTRATO DE CONTA CORRENTE", .xlsx): Data por extenso ("Quinta, 30 de abril de
  2026") / Descrição / Crédito (R$) / Débito (R$); SEM saldo, sem documento e sem o CPF do Pix; mês inteiro,
  sem o corte de 440 linhas. Só serve para COMPLETAR buracos provados pelo saldo (ver servico.completar_buracos).
Diferença para a Gestão: lê também os débitos (lá só interessam créditos) e confere o saldo corrido.
Linhas "SALDO ANTERIOR" e lançamentos de valor zero (ex.: "ABERTURA") não viram lançamento.
"""
import csv
import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

CENTAVO = Decimal("0.01")


@dataclass
class Linha:
    data: date
    historico: str
    documento: str | None
    valor: Decimal            # com sinal: saída < 0
    saldo: Decimal | None


@dataclass
class Extrato:
    banco: str
    agencia: str | None
    conta: str | None
    formato: str
    linhas: list[Linha] = field(default_factory=list)   # ordem cronológica
    saldo_inicial: Decimal | None = None
    quebras_saldo: int = 0

    @property
    def saldo_final(self) -> Decimal | None:
        return next((l.saldo for l in reversed(self.linhas) if l.saldo is not None), None)


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def _num(v) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(CENTAVO)
    s = str(v).strip().replace("R$", "").strip()
    if not s or s.lower() == "nan":
        return None
    if re.fullmatch(r"-?\d+(\.\d+)?", s):          # "1234.5" (número já sem milhar)
        return Decimal(s).quantize(CENTAVO)
    try:
        return Decimal(s.replace(".", "").replace(",", ".")).quantize(CENTAVO)
    except InvalidOperation:
        return None


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:10]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _linhas_do_arquivo(caminho: Path) -> tuple[list[list], str]:
    suf = caminho.suffix.lower()
    if suf == ".csv":
        texto = caminho.read_bytes().decode("latin1")
        # o Bradesco exporta com quebras de linha mistas (CR CR LF): splitlines(), como na Gestão
        return [r for r in csv.reader(texto.splitlines(), delimiter=";")], "csv"
    if suf == ".xls":
        import xlrd
        ws = xlrd.open_workbook(str(caminho)).sheet_by_index(0)
        return [ws.row_values(i) for i in range(ws.nrows)], "xls"
    if suf == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(caminho), read_only=True, data_only=True)
        for ws in wb.worksheets:                     # 1ª aba que tiver o cabeçalho de extrato
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
            if _cabecalho(linhas) is not None:
                return linhas, "xlsx"
        return [], "xlsx"
    raise ValueError(f"Formato não suportado: {caminho.name}")


_MESES = {"janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7,
          "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12}


def _data_por_extenso(v) -> date | None:
    """'Quinta, 30 de abril de 2026' → date(2026, 4, 30)."""
    m = re.search(r"(\d{1,2}) de (\w+) de (\d{4})", str(v or ""))
    if not m or m.group(2).lower() not in _MESES:
        return None
    return date(int(m.group(3)), _MESES[m.group(2).lower()], int(m.group(1)))


def _cabecalho(linhas: list[list]) -> int | None:
    for i, r in enumerate(linhas[:20]):
        vals = [str(v).strip() for v in r if v is not None]
        if "Data" in vals and ("Lançamento" in vals or "Histórico" in vals or "Descrição" in vals):
            return i
    return None


def ler(caminho: Path) -> Extrato:
    linhas, tipo = _linhas_do_arquivo(caminho)
    h = _cabecalho(linhas)
    if h is None:
        raise ValueError(f"{caminho.name}: não achei o cabeçalho de extrato (Data / Lançamento ou Histórico)")
    cab = [str(v).strip() if v is not None else "" for v in linhas[h]]
    topo = " ".join(str(v) for r in linhas[:h] for v in r if v not in (None, ""))
    if {"Movimentação", "Descrição", "Classificação"} <= set(cab):
        raise ValueError(f"{caminho.name}: é planilha TRATADA (sem saldo); serve para classificação, não como extrato")

    if "Descrição" in cab and "Crédito (R$)" in cab and "Saldo (R$)" not in cab:   # Santander internet banking
        m = re.search(r"Conta:\s*([\d-]+)", topo)
        ext = Extrato("Santander", None, m.group(1).replace("-", "") if m else None, f"santander_ib_{tipo}")
        brutas = []
        for r in linhas[h + 1:]:
            d = dict(zip(cab, list(r) + [None] * (len(cab) - len(r))))
            dia = _data_por_extenso(d.get("Data"))
            if dia is None:
                continue
            valor = (_num(d.get("Crédito (R$)")) or Decimal(0)) - abs(_num(d.get("Débito (R$)")) or Decimal(0))
            if valor:
                brutas.append(Linha(dia, " ".join(str(d.get("Descrição") or "").split()), None, valor, None))
    elif "Crédito (R$)" in cab:                                   # Bradesco
        m = re.search(r"Ag[êe]ncia:\s*(\d+)\s+Conta:\s*([\d-]+)", topo)
        ext = Extrato("Bradesco", m.group(1) if m else None, m.group(2) if m else None, f"bradesco_{tipo}")
        brutas = []
        for r in linhas[h + 1:]:
            d = dict(zip(cab, list(r) + [None] * (len(cab) - len(r))))
            if str(d.get("Data") or "").strip() == "Total":
                break
            dia = _data(d.get("Data"))
            if dia is None:
                continue
            hist = str(d.get("Lançamento") or "").strip()
            saldo = _num(d.get("Saldo (R$)"))
            if hist.upper().startswith("SALDO ANTERIOR"):
                ext.saldo_inicial = saldo
                continue
            valor = (_num(d.get("Crédito (R$)")) or Decimal(0)) - abs(_num(d.get("Débito (R$)")) or Decimal(0))
            if valor == 0:
                continue
            doc = str(d.get("Dcto.") or "").strip() or None
            brutas.append(Linha(dia, hist, doc.removesuffix(".0") if doc else None, valor, saldo))
    elif "Valor (R$)" in cab:                                     # Santander
        m = re.search(r"AGENCIA\s+(\d+)\s+CONTA\s+(\d+)", topo.upper())
        ext = Extrato("Santander", m.group(1) if m else None, m.group(2) if m else None, f"santander_{tipo}")
        brutas = []
        for r in linhas[h + 1:]:
            d = dict(zip(cab, list(r) + [None] * (len(cab) - len(r))))
            dia = _data(d.get("Data"))
            valor = _num(d.get("Valor (R$)"))
            if dia is None or not valor:
                continue
            doc = str(d.get("Documento") or "").strip() or None
            brutas.append(Linha(dia, " ".join(str(d.get("Histórico") or "").split()), doc,
                                valor, _num(d.get("Saldo (R$)"))))
    else:
        raise ValueError(f"{caminho.name}: cabeçalho não reconhecido ({', '.join(c for c in cab if c)})")

    if len(brutas) > 1 and brutas[0].data > brutas[-1].data:     # mais novo primeiro → inverte
        brutas.reverse()
    ext.linhas = brutas
    anterior = ext.saldo_inicial
    for l in brutas:
        if anterior is not None and l.saldo is not None and abs(anterior + l.valor - l.saldo) > CENTAVO:
            ext.quebras_saldo += 1
        if l.saldo is not None:
            anterior = l.saldo
    if ext.saldo_inicial is None and brutas and brutas[0].saldo is not None:
        ext.saldo_inicial = brutas[0].saldo - brutas[0].valor
    return ext
