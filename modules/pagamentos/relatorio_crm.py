"""Relatório CRM_Fornecedores_<Cliente>.xlsx gerado A PARTIR do banco (não é fonte paralela).

    python -m modules.pagamentos.relatorio_crm --tenant rodiza            # grava em /data/dashboards/
    GET /tenants/{id}/fornecedores/relatorio.xlsx                          # download pela API

Formato combinado em 28/09: 2 abas, título mesclado, subtítulo em itálico com a origem dos dados,
cabeçalho escuro com fonte branca em negrito, texto com quebra, larguras ajustadas, painel congelado.
"""
import argparse
import io
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from modules.pagamentos.models import DespesaRecorrente, Fornecedor, PadraoPagamento

NOMES_PADRAO = {
    "um_para_um": "1 nota = 1 débito",
    "consolidado_mensal": "Consolidado no mês",
    "agrupado_vencimento": "Agrupado por vencimento",
    "nota_mais_ajuste": "Nota + ajuste",
    "livre": "Livre",
}
_CABECALHO = PatternFill("solid", fgColor="1F2937")
_QUEBRA = Alignment(wrap_text=True, vertical="top")


def _moeda(v: Decimal) -> str:
    inteiro = v == v.to_integral_value()
    texto = f"{v:,.0f}" if inteiro else f"{v:,.2f}"
    return "R$ " + texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _cnpj(c: str | None) -> str | None:
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if c else None


def _dias(r: DespesaRecorrente) -> str:
    return f"{r.dia_inicio:02d}" if r.dia_inicio == r.dia_fim else f"{r.dia_inicio:02d}-{r.dia_fim:02d}"


def _valor(r: DespesaRecorrente) -> str:
    if r.valor_variavel or r.valor_min is None:
        return "variável"
    if r.valor_max is None or r.valor_max == r.valor_min:
        return _moeda(r.valor_min)
    return f"{_moeda(r.valor_min)} – {_moeda(r.valor_max)[3:]}"


def _aba(ws, titulo: str, subtitulo: str, colunas: list[tuple[str, int]], linhas: list[list]) -> None:
    n = len(colunas)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.cell(1, 1, titulo).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.cell(2, 1, subtitulo).font = Font(italic=True, size=9, color="4B5563")
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    for i, (nome, _) in enumerate(colunas, start=1):
        c = ws.cell(3, i, nome)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), _CABECALHO, _QUEBRA
    for r, linha in enumerate(linhas, start=4):
        for i, valor in enumerate(linha, start=1):
            ws.cell(r, i, valor).alignment = _QUEBRA
    for i, (nome, largura_max) in enumerate(colunas, start=1):
        maior = max([len(nome)] + [len(str(l[i - 1])) for l in linhas if l[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(maior + 2, 10), largura_max)
    ws.freeze_panes = "A4"


def gerar(db: Session, tenant_id: int, cliente: str) -> bytes:
    fornecedores = db.scalars(select(Fornecedor).where(Fornecedor.tenant_id == tenant_id, Fornecedor.ativo.is_(True))
                              .order_by(Fornecedor.nome)).all()
    padroes = {p.fornecedor_id: p for p in db.scalars(select(PadraoPagamento).where(
        PadraoPagamento.tenant_id == tenant_id, PadraoPagamento.vigencia_fim.is_(None)))}
    recorrentes = db.scalars(select(DespesaRecorrente).where(DespesaRecorrente.tenant_id == tenant_id)
                             .order_by(DespesaRecorrente.dia_inicio, DespesaRecorrente.descricao)).all()

    agora = datetime.now().strftime("%d-%m-%Y %H:%M")
    subtitulo = (f"Gerado do IN-360 em {agora}. Carga inicial: auditoria caixa/banco/DRE (Jan–Set/2026) e e-mail "
                 "rodizaalimentos@gmail.com. CNPJ/endereço em branco precisam ser confirmados antes de qualquer "
                 "uso em contrato.")
    wb = Workbook()
    ws = wb.active
    ws.title = "CRM Fornecedores"
    _aba(ws, f"CRM de Fornecedores — {cliente}", subtitulo,
         [("Fornecedor (uso no dia a dia)", 28), ("Também aparece como", 24), ("Razão Social", 40), ("CNPJ", 20),
          ("Endereço", 30), ("Categoria", 34), ("Prazo de Pagamento", 40), ("Padrão de pagamento", 22),
          ("Observações", 60)],
         [[f.nome, ", ".join(f.apelidos) or None, f.razao_social, _cnpj(f.cnpj), f.endereco, f.categoria,
           f.prazo_pagamento, NOMES_PADRAO.get(padroes[f.id].tipo) if f.id in padroes else None, f.observacoes]
          for f in fornecedores])

    ws2 = wb.create_sheet("Despesas Fixas")
    _aba(ws2, f"Despesas Fixas — {cliente}", subtitulo,
         [("Dia do mês", 10), ("Fornecedor", 44), ("Valor aproximado", 24), ("Situação", 22), ("Observação", 60)],
         [[_dias(r), r.descricao, _valor(r),
           "Ativa" if r.ativo else (f"Encerrada em {r.vigencia_fim:%d-%m-%Y}" if r.vigencia_fim else "Inativa"),
           r.observacao] for r in recorrentes])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    import modules.models  # noqa: F401
    from core.db import nova_sessao, usar_tenant
    from modules.tenants.models import Tenant

    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(prog="python -m modules.pagamentos.relatorio_crm")
    ap.add_argument("--tenant", required=True)
    ap.add_argument("--pasta", default="/data/dashboards", type=Path)
    a = ap.parse_args()
    with nova_sessao() as db:
        t = db.scalar(select(Tenant).where(Tenant.codigo == a.tenant))
        if t is None:
            sys.exit(f"Cliente '{a.tenant}' não existe.")
        usar_tenant(db, t.id)
        conteudo = gerar(db, t.id, a.tenant.capitalize())
    destino = a.pasta / f"CRM_Fornecedores_{a.tenant.capitalize()}.xlsx"
    destino.write_bytes(conteudo)
    print(f"Gerado: {destino} ({len(conteudo):,} bytes)")


if __name__ == "__main__":
    main()
