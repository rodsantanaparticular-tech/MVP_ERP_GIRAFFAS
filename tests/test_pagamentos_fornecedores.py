"""Fase 1B, passo 1 — CRM de fornecedores, padrões de pagamento e despesas recorrentes.

Roda contra o Postgres real (papel da aplicação, com RLS), na stack de TESTE:
    infra/testes/rodar_testes.sh
"""
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

import modules.models  # noqa: F401
from api.main import app
from core.db import nova_sessao, usar_tenant
from modules.pagamentos.carga import carregar
from modules.pagamentos.models import DespesaRecorrente, Fornecedor, PadraoPagamento
from tests.test_etapa1_seguranca_cofre import _sufixo, criar_usuario, login

client = TestClient(app)
SEED = Path(__file__).resolve().parents[1] / "database" / "seeds" / "fornecedores_rodiza.json"


@pytest.fixture(scope="module")
def cen():
    adm = login(criar_usuario("plataforma_admin"))
    ids = {}
    for nome in ("A", "B"):
        r = client.post("/tenants", headers=adm, json={"rede_codigo": "giraffas", "codigo": f"p{nome}{_sufixo()}",
                                                         "razao_social": f"Cliente {nome}"})
        assert r.status_code == 201, r.text
        tid = r.json()["id"]
        loja = client.post(f"/tenants/{tid}/lojas", headers=adm, json={"nome": "Loja", "filial_sigla": "L1"}).json()
        ids[nome] = {"tenant": tid, "loja": loja["id"], "codigo": r.json()["codigo"]}
    return {"adm": adm, **ids}


def _novo_fornecedor(cen, tenant="A", **extra) -> dict:
    corpo = {"nome": f"Forn {_sufixo()}", **extra}
    r = client.post(f"/tenants/{cen[tenant]['tenant']}/fornecedores", headers=cen["adm"], json=corpo)
    assert r.status_code == 201, r.text
    return r.json()


def test_cria_fornecedor_normaliza_cnpj_e_bloqueia_duplicado(cen):
    f = _novo_fornecedor(cen, cnpj="49.319.411/0021-87", apelidos=["MB"])
    assert f["cnpj"] == "49319411002187" and f["apelidos"] == ["MB"] and f["padrao_vigente"] is None
    tid = cen["A"]["tenant"]
    r = client.post(f"/tenants/{tid}/fornecedores", headers=cen["adm"], json={"nome": f["nome"]})
    assert r.status_code == 409
    r = client.post(f"/tenants/{tid}/fornecedores", headers=cen["adm"],
                    json={"nome": "Outro nome", "cnpj": "49319411002187"})
    assert r.status_code == 409
    r = client.post(f"/tenants/{tid}/fornecedores", headers=cen["adm"], json={"nome": "X", "cnpj": "123"})
    assert r.status_code == 422


def test_trocar_padrao_preserva_historico(cen):
    tid = cen["A"]["tenant"]
    f = _novo_fornecedor(cen)
    url = f"/tenants/{tid}/fornecedores/{f['id']}"
    r = client.put(f"{url}/padrao", headers=cen["adm"], json={"tipo": "um_para_um", "vigencia_inicio": "2026-01-01"})
    assert r.status_code == 200, r.text
    # mesmo dia = correção: substitui, sem criar histórico
    r = client.put(f"{url}/padrao", headers=cen["adm"],
                   json={"tipo": "um_para_um", "vigencia_inicio": "2026-01-01", "tolerancia_valor": "0.05"})
    assert r.status_code == 200 and float(r.json()["tolerancia_valor"]) == 0.05
    # novo padrão: encerra o anterior na véspera
    r = client.put(f"{url}/padrao", headers=cen["adm"],
                   json={"tipo": "consolidado_mensal", "vigencia_inicio": "2026-09-01"})
    assert r.status_code == 200, r.text
    hist = client.get(f"{url}/padroes", headers=cen["adm"]).json()
    assert [(p["tipo"], p["vigencia_inicio"], p["vigencia_fim"]) for p in hist] == [
        ("consolidado_mensal", "2026-09-01", None), ("um_para_um", "2026-01-01", "2026-08-31")]
    # não pode começar antes do vigente
    r = client.put(f"{url}/padrao", headers=cen["adm"], json={"tipo": "livre", "vigencia_inicio": "2026-05-01"})
    assert r.status_code == 422
    lista = client.get(f"/tenants/{tid}/fornecedores", headers=cen["adm"]).json()
    assert next(x for x in lista if x["id"] == f["id"])["padrao_vigente"]["tipo"] == "consolidado_mensal"


def test_desativar_some_da_lista(cen):
    tid = cen["A"]["tenant"]
    f = _novo_fornecedor(cen)
    r = client.patch(f"/tenants/{tid}/fornecedores/{f['id']}", headers=cen["adm"], json={"ativo": False})
    assert r.status_code == 200 and r.json()["ativo"] is False
    ativos = client.get(f"/tenants/{tid}/fornecedores", headers=cen["adm"]).json()
    todos = client.get(f"/tenants/{tid}/fornecedores?incluir_inativos=true", headers=cen["adm"]).json()
    assert f["id"] not in {x["id"] for x in ativos} and f["id"] in {x["id"] for x in todos}


def test_perfis_ler_x_alterar(cen):
    tid, loja = cen["A"]["tenant"], cen["A"]["loja"]
    f = _novo_fornecedor(cen)
    gestor = login(criar_usuario("gestor_loja", tid, loja))       # conciliacao:ler, sem parametros:gerir
    assert client.get(f"/tenants/{tid}/fornecedores", headers=gestor).status_code == 200
    assert client.get(f"/tenants/{cen['B']['tenant']}/fornecedores", headers=gestor).status_code == 403
    assert client.post(f"/tenants/{tid}/fornecedores", headers=gestor, json={"nome": "Z"}).status_code == 403
    assert client.put(f"/tenants/{tid}/fornecedores/{f['id']}/padrao", headers=gestor,
                      json={"tipo": "livre"}).status_code == 403
    ti = login(criar_usuario("ti_ops", tid))                         # sem conciliacao:ler
    assert client.get(f"/tenants/{tid}/fornecedores", headers=ti).status_code == 403
    eng = login(criar_usuario("eng_dados", tid))                     # parametros:gerir
    assert client.post(f"/tenants/{tid}/fornecedores", headers=eng,
                       json={"nome": f"Eng {_sufixo()}"}).status_code == 201


def test_isolamento_entre_clientes(cen):
    fa = _novo_fornecedor(cen, "A")
    tb = cen["B"]["tenant"]
    eng_b = login(criar_usuario("eng_dados", tb))
    assert client.get(f"/tenants/{cen['A']['tenant']}/fornecedores", headers=eng_b).status_code == 403
    # pela rota do próprio cliente B, o fornecedor de A não existe
    assert client.patch(f"/tenants/{tb}/fornecedores/{fa['id']}", headers=eng_b,
                        json={"nome": "roubado"}).status_code == 404
    # recorrente de B apontando para fornecedor de A
    r = client.post(f"/tenants/{tb}/despesas-recorrentes", headers=eng_b,
                    json={"descricao": "x", "fornecedor_id": fa["id"], "dia_inicio": 1, "dia_fim": 2})
    assert r.status_code == 404
    # no banco (RLS): com o tenant B ativo, as linhas de A são invisíveis
    with nova_sessao() as db:
        usar_tenant(db, tb)
        assert db.get(Fornecedor, fa["id"]) is None
        assert db.scalar(select(func.count()).select_from(Fornecedor)
                         .where(Fornecedor.tenant_id == cen["A"]["tenant"])) == 0


def test_recorrente_valida_dias_e_valores(cen):
    tid = cen["A"]["tenant"]
    base = {"descricao": f"Fixa {_sufixo()}", "dia_inicio": 10, "dia_fim": 12, "valor_min": "50", "valor_max": "150"}
    assert client.post(f"/tenants/{tid}/despesas-recorrentes", headers=cen["adm"],
                       json={**base, "dia_inicio": 15}).status_code == 422
    assert client.post(f"/tenants/{tid}/despesas-recorrentes", headers=cen["adm"],
                       json={**base, "valor_min": "200"}).status_code == 422
    r = client.post(f"/tenants/{tid}/despesas-recorrentes", headers=cen["adm"], json=base)
    assert r.status_code == 201, r.text
    assert client.post(f"/tenants/{tid}/despesas-recorrentes", headers=cen["adm"], json=base).status_code == 409


def test_carga_do_crm_real_e_idempotente(cen):
    codigo, tid = cen["B"]["codigo"], cen["B"]["tenant"]
    relato = carregar(codigo, SEED, simular=True)
    assert relato[-1].strip().startswith("SIMULAÇÃO")
    with nova_sessao() as db:
        usar_tenant(db, tid)
        assert db.scalar(select(func.count()).select_from(Fornecedor)) == 0

    relato = carregar(codigo, SEED, simular=False)
    assert sum(l.startswith("[criado ] fornecedor") for l in relato) == 14
    with nova_sessao() as db:
        usar_tenant(db, tid)
        forn = {f.nome: f for f in db.scalars(select(Fornecedor))}
        assert len(forn) == 14 and forn["RFG"].apelidos == ["MB"] and forn["RFG"].cnpj == "49319411002187"
        padroes = {p.fornecedor_id: p for p in db.scalars(select(PadraoPagamento))}
        assert len(padroes) == 14
        assert padroes[forn["Kelminas"].id].tipo == "consolidado_mensal"
        assert padroes[forn["JM"].id].tipo == "nota_mais_ajuste"
        assert padroes[forn["Ribeiro"].id].vigencia_inicio == date(2026, 1, 1)
        rec = {r.descricao: r for r in db.scalars(select(DespesaRecorrente))}
        assert len(rec) == 9
        taxa = rec["Multiplan (Taxa de Transferência)"]
        assert taxa.ativo is False and taxa.vigencia_fim == date(2026, 6, 20)
        assert rec["Cartão Crédito Anuidade"].fornecedor_id is None
        assert rec["Giraffas Administradora de Franquia (royalties)"].valor_variavel is True
        assert rec["GCOM (sistema)"].fornecedor_id == forn["GCOM"].id

    # segunda rodada: nada novo
    relato = carregar(codigo, SEED, simular=False)
    assert not any(l.startswith("[criado ") for l in relato)
    # a API mostra 8 recorrentes ativas (a taxa encerrada fica de fora)
    assert len(client.get(f"/tenants/{tid}/despesas-recorrentes", headers=cen["adm"]).json()) == 8


def test_relatorio_excel_do_crm(cen):
    import io

    from openpyxl import load_workbook
    tid = cen["B"]["tenant"]   # já carregado pelo teste da carga (roda antes, mesma ordem do arquivo)
    r = client.get(f"/tenants/{tid}/fornecedores/relatorio.xlsx", headers=cen["adm"])
    assert r.status_code == 200 and "CRM_Fornecedores_" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["CRM Fornecedores", "Despesas Fixas"]
    crm, fixas = wb["CRM Fornecedores"], wb["Despesas Fixas"]
    assert crm.freeze_panes == "A4" and crm["A3"].font.bold and crm["A2"].font.italic
    linhas = {row[0]: row for row in crm.iter_rows(min_row=4, values_only=True)}
    assert len(linhas) == 14
    assert linhas["RFG"][1] == "MB" and linhas["RFG"][3] == "49.319.411/0021-87"
    assert linhas["Kelminas"][7] == "Consolidado no mês"
    fx = {row[1]: row for row in fixas.iter_rows(min_row=4, values_only=True)}
    assert len(fx) == 9
    assert fx["Multiplan (CAM + aluguel, referente ao mês anterior)"][2] == "R$ 38.000 – 44.000"
    assert fx["GCOM (sistema)"][2] == "R$ 552,94" and fx["GCOM (sistema)"][0] == "05-08"
    assert fx["Multiplan (Taxa de Transferência)"][3] == "Encerrada em 20-06-2026"
    # gestor de loja (só leitura) também baixa; cliente alheio não
    gestor = login(criar_usuario("gestor_loja", cen["A"]["tenant"], cen["A"]["loja"]))
    assert client.get(f"/tenants/{cen['A']['tenant']}/fornecedores/relatorio.xlsx", headers=gestor).status_code == 200
    assert client.get(f"/tenants/{tid}/fornecedores/relatorio.xlsx", headers=gestor).status_code == 403
